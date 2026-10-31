"""Write tracker.xlsx, charts, snapshot JSON and summary.md."""
from __future__ import annotations

import datetime as dt
import json

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import analysis
from .config import ROOT

# Reference palette (validated categorical order) and text tokens
SERIES = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300", "#6250d6", "#e34948"]
BUCKET_COLOR = dict(zip(analysis.BUCKET_ORDER, SERIES[:4] + [SERIES[5]]))  # slot 5 (magenta) is kept for macro lines
BUCKET_MARKER = {"Bottleneck": "o", "Compute leader": "s", "Hyperscaler": "D", "Leveraged buildout": "^", "Adjacent": "P"}
MACRO_COLOR = [SERIES[6], SERIES[4]]   # violet, magenta: never used for groups
INK, INK2, GRID, SURFACE = "#0b0b0b", "#52514e", "#e4e3df", "#fcfcfb"
NEG_FILL, POS_FILL, MID_FILL = "F4B6B5", "B9DFC6", "FFFFFF"

PCT_HINTS = ("ret_", "pct_", "yoy", "margin", "upside", "revision", "short_pct", "chg_4q", "chg_2q", "_vs_3y",
             "capex_pct", "debt_chg", "price_chg", "move", "Bottleneck", "Compute", "Hyperscaler", "Leveraged",
             "SoftBank")
SCALE_HINTS = ("ret_", "score", "z_", "price_chg", "pct_from", "pct_vs", "revenue_yoy", "gap", "change", "news",
               "Bottleneck", "Compute", "Hyperscaler", "Leveraged", "SoftBank", "upside", "revision")


def _style_plt():
    plt.rcParams.update({"font.size": 10, "axes.edgecolor": GRID, "axes.labelcolor": INK2, "xtick.color": INK2,
                         "ytick.color": INK2, "axes.grid": True, "grid.color": GRID, "grid.linewidth": 0.6,
                         "axes.spines.top": False, "axes.spines.right": False, "figure.facecolor": SURFACE,
                         "axes.facecolor": SURFACE, "axes.titlecolor": INK, "axes.titleweight": "bold",
                         "axes.titlesize": 11, "legend.frameon": False, "lines.linewidth": 2})


def _footer(fig, text):
    fig.text(0.01, 0.005, text, fontsize=7.5, color=INK2, ha="left", va="bottom")


def charts(ds, sc: pd.DataFrame, out_dir, asof: str) -> list[str]:
    _style_plt()
    out_dir.mkdir(exist_ok=True)
    paths = []

    # 1. Group indexes in USD and THB
    fig, axes = plt.subplots(1, 2, figsize=(12, 4.6), sharey=False)
    for ax, ccy in zip(axes, ("USD", "THB")):
        idx = analysis.group_index(ds, ccy)
        for b in idx:
            ax.plot(idx.index, idx[b], color=BUCKET_COLOR[b], label=b)
            ax.annotate(f"{idx[b].iloc[-1]:.0f}", (idx.index[-1], idx[b].iloc[-1]), xytext=(4, 0),
                        textcoords="offset points", fontsize=8, color=INK2, va="center")
        ax.axhline(100, color=INK2, lw=0.8, ls=":")
        ax.set_title(f"Group index, {ccy} (start = 100)")
    axes[0].legend(loc="upper left", fontsize=8)
    _footer(fig, f"Equal-weight buy-and-hold in USD; rebalanced only when a member lists (SNDK 2025-02, CRWV 2025-03). "
                 f"Source: yfinance closes, Yahoo FX. As of {asof}.")
    fig.tight_layout(rect=(0, 0.03, 1, 1))
    p = out_dir / "group_index_usd_thb.png"; fig.savefig(p, dpi=150); plt.close(fig); paths.append(p)

    # 2. Sentiment vs fundamentals
    fig, ax = plt.subplots(figsize=(8, 7))
    d = sc.dropna(subset=["sentiment_score", "fundamentals_score"])
    for b in analysis.BUCKET_ORDER:
        g = d[d.bucket == b]
        ax.scatter(g.fundamentals_score, g.sentiment_score, s=70, color=BUCKET_COLOR[b], marker=BUCKET_MARKER[b],
                   edgecolor=SURFACE, linewidth=2, label=b, zorder=3)
    texts = [ax.text(r.fundamentals_score, r.sentiment_score, tk, fontsize=8, color=INK) for tk, r in d.iterrows()]
    ax.axhline(0, color=INK2, lw=0.8); ax.axvline(0, color=INK2, lw=0.8)
    lim = max(1.5, np.nanmax(np.abs(d[["sentiment_score", "fundamentals_score"]].values)) * 1.15) if len(d) else 2
    ax.set_xlim(-lim, lim); ax.set_ylim(-lim, lim)
    for x, y, t in ((lim * .97, -lim * .97, "good fundamentals,\nweak sentiment"),
                    (-lim * .97, lim * .97, "weak fundamentals,\nstrong sentiment")):
        ax.text(x, y, t, ha="right" if x > 0 else "left", va="bottom" if y < 0 else "top", fontsize=8, color=INK2)
    ax.set_xlabel("Fundamentals-momentum score (z)"); ax.set_ylabel("Sentiment score (z)")
    ax.set_title("Sentiment vs fundamentals")
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.09), ncol=4, fontsize=8)
    try:
        from adjustText import adjust_text
        adjust_text(texts, ax=ax, expand=(1.3, 1.6), arrowprops=dict(arrowstyle="-", color=GRID, lw=0.6))
    except ImportError:
        pass
    _footer(fig, f"z-score blends across the universe. Source: yfinance, SEC, Google News. As of {asof}.")
    fig.tight_layout(rect=(0, 0.03, 1, 1))
    p = out_dir / "sentiment_vs_fundamentals.png"; fig.savefig(p, dpi=150); plt.close(fig); paths.append(p)

    # 3. Yields (top) and leveraged group (bottom): two panels, never a dual axis
    m = analysis.window(analysis.macro_frame(ds))
    lev = analysis.group_index(ds, "USD")
    fig, (a1, a2) = plt.subplots(2, 1, figsize=(11, 6.5), sharex=True, gridspec_kw={"height_ratios": [1, 1]})
    for k, col in (("US_10Y", MACRO_COLOR[0]), ("JGB_30Y", MACRO_COLOR[1])):
        if k in m:
            s = m[k].dropna()
            a1.plot(s.index, s, color=col, label={"US_10Y": "US 10-year", "JGB_30Y": "JGB 30-year"}[k])
            a1.annotate(f"{s.iloc[-1]:.2f}%", (s.index[-1], s.iloc[-1]), xytext=(4, 0), textcoords="offset points",
                        fontsize=8, color=INK2, va="center")
    a1.set_title("Yields (%)"); a1.legend(loc="upper left", fontsize=8)
    if "Leveraged buildout" in lev:
        a2.plot(lev.index, lev["Leveraged buildout"], color=BUCKET_COLOR["Leveraged buildout"], label="Leveraged buildout")
        if "Bottleneck" in lev:
            a2.plot(lev.index, lev["Bottleneck"], color=BUCKET_COLOR["Bottleneck"], label="Bottleneck")
        a2.axhline(100, color=INK2, lw=0.8, ls=":")
        a2.legend(loc="upper left", fontsize=8)
    a2.set_title("Group index, USD (start = 100)")
    _footer(fig, f"Source: FRED DGS10; Japan MoF JGB yield CSV; yfinance. As of {asof}.")
    fig.tight_layout(rect=(0, 0.03, 1, 1))
    p = out_dir / "yields_vs_leveraged.png"; fig.savefig(p, dpi=150); plt.close(fig); paths.append(p)

    # 4. USD/JPY (top, its own unit) vs SoftBank and the leveraged group (bottom, indexed)
    fig, (a1, a2) = plt.subplots(2, 1, figsize=(11, 6.5), sharex=True)
    start = lev.index[0] if len(lev) else None
    if "USDJPY" in m:
        s = m["USDJPY"].dropna()
        s = s[s.index >= start] if start is not None else s
        a1.plot(s.index, s, color=MACRO_COLOR[0], label="USD/JPY")
        a1.annotate(f"{s.iloc[-1]:.1f}", (s.index[-1], s.iloc[-1]), xytext=(4, 0), textcoords="offset points",
                    fontsize=8, color=INK2, va="center")
    a1.set_title("USD/JPY (yen per dollar; falling = yen stronger)")
    for lab, s, col in (("SoftBank 9984.T (JPY)", ds.prices.get("9984.T"), MACRO_COLOR[1]),
                        ("Leveraged buildout group (USD)", lev.get("Leveraged buildout"), BUCKET_COLOR["Leveraged buildout"])):
        if s is None:
            continue
        s = s[s.index >= start].dropna() if start is not None else s.dropna()
        s = s / s.iloc[0] * 100
        a2.plot(s.index, s, color=col, label=lab)
        a2.annotate(f"{s.iloc[-1]:.0f}", (s.index[-1], s.iloc[-1]), xytext=(4, 0), textcoords="offset points",
                    fontsize=8, color=INK2, va="center")
    a2.axhline(100, color=INK2, lw=0.8, ls=":")
    a2.set_title("SoftBank and the leveraged group (start = 100)"); a2.legend(loc="upper left", fontsize=8)
    _footer(fig, f"Source: Yahoo JPY=X, yfinance. As of {asof}.")
    fig.tight_layout(rect=(0, 0.03, 1, 1))
    p = out_dir / "usdjpy_vs_softbank.png"; fig.savefig(p, dpi=150); plt.close(fig); paths.append(p)
    return paths


# ---------------- Excel ----------------

def _write_tab(writer, name, df, title, source, asof, index=False):
    df = df.copy()
    for c in df.columns:
        if df[c].dtype == object:
            df[c] = df[c].map(lambda v: v.isoformat() if isinstance(v, (dt.date, pd.Timestamp)) else v)
    df.to_excel(writer, sheet_name=name, startrow=3, index=index)
    ws = writer.sheets[name]
    ws["A1"] = title; ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"Source: {source}"; ws["A3"] = f"As of: {asof}. Research only, not trading advice."
    for c in ("A2", "A3"):
        ws[c].font = Font(italic=True, color="52514E", size=9)
    ws.freeze_panes = ws.cell(row=5, column=2 if index or len(df.columns) > 3 else 1)
    hdr = 4
    ncols = len(df.columns) + (df.index.nlevels if index else 0)
    for j in range(1, ncols + 1):
        cell = ws.cell(row=hdr, column=j)
        cell.font = Font(bold=True); cell.fill = PatternFill("solid", fgColor="EDECE8")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    names = ([df.index.name or ""] * df.index.nlevels if index else []) + [str(c) for c in df.columns]
    last = hdr + len(df)
    for j, nm in enumerate(names, 1):
        L = get_column_letter(j)
        ws.column_dimensions[L].width = min(48, max(10, len(nm) + 2))
        if len(df) == 0:
            continue
        col = df[nm] if nm in df.columns else None
        numeric = col is not None and pd.api.types.is_numeric_dtype(col)
        rng = f"{L}{hdr + 1}:{L}{last}"
        if numeric and any(h in nm for h in PCT_HINTS) and not nm.startswith("n_"):
            for row in ws[rng]:
                for cell in row:
                    cell.number_format = "0.0%"
        elif numeric:
            for row in ws[rng]:
                for cell in row:
                    cell.number_format = "#,##0.00"
        if numeric and any(h in nm for h in SCALE_HINTS) and not nm.startswith("n_"):
            ws.conditional_formatting.add(rng, ColorScaleRule(start_type="percentile", start_value=5, start_color=NEG_FILL,
                                                              mid_type="num", mid_value=0, mid_color=MID_FILL,
                                                              end_type="percentile", end_value=95, end_color=POS_FILL))
        if nm in ("status", "verdict", "live"):
            for v, color in (("pass", POS_FILL), ("fail", NEG_FILL), ("Supported", POS_FILL), ("Not supported", NEG_FILL),
                             ("matches", POS_FILL), ("differs", NEG_FILL), ("True", "FCE7B2")):
                ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{v}"'],
                                                              fill=PatternFill("solid", fgColor=color)))
    return ws


def workbook(path, tabs: list[tuple]):
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        for t in tabs:
            _write_tab(w, *t)


def scorecard_frame(sc_results) -> pd.DataFrame:
    rows = []
    for h in sc_results:
        rows.append({"hypothesis": h["id"], "item": "VERDICT", "name": h["name"], "status": h["verdict"],
                     "value": None, "op": "", "threshold": None, "detail": h["why"],
                     "previous": h["prev_verdict"] or "(no earlier snapshot)"})
        for t in h["tests"]:
            rows.append({"hypothesis": h["id"], "item": t["id"], "name": t["name"], "status": t["status"],
                         "value": t.get("value"), "op": t.get("op", ""), "threshold": t.get("threshold"),
                         "detail": t["detail"], "previous": ""})
            for i, s in enumerate(t.get("subs", []), 1):
                rows.append({"hypothesis": h["id"], "item": f"{t['id']}.{i}", "name": "  " + s.get("label", ""),
                             "status": s["status"], "value": s.get("value"), "op": s.get("op", ""),
                             "threshold": s.get("threshold"), "detail": s["detail"], "previous": ""})
        for s in h["signals"]:
            rows.append({"hypothesis": h["id"], "item": f"{s['id']} ({s['kind']})", "name": s["name"],
                         "status": "LIVE" if s["live"] else s["status"], "value": s.get("value"),
                         "op": s.get("op", ""), "threshold": s.get("threshold"), "detail": s["detail"], "previous": ""})
    return pd.DataFrame(rows)


# ---------------- snapshot & summary ----------------

def snapshot(path, asof, sc_results, ct, scores, macro_changes):
    keep = ["price", "ret_1m", "ret_3m", "ret_ytd", "ret_1y", "pct_from_52w_high", "forward_pe", "ev_ebitda_ttm",
            "revenue_yoy", "operating_margin", "net_debt_to_ebitda", "analyst_rating_mean", "target_upside",
            "eps_revision_90d", "short_pct_float", "news_sentiment", "corr90_us10y"]
    comp = {}
    for tk, r in ct.iterrows():
        comp[tk] = {k: (None if pd.isna(r.get(k)) else float(r.get(k))) for k in keep if k in r}
        for k in ("sentiment_score", "fundamentals_score"):
            v = scores.loc[tk, k] if tk in scores.index else None
            comp[tk][k] = None if v is None or pd.isna(v) else float(v)
    js = {"date": asof,
          "hypotheses": [{"id": h["id"], "verdict": h["verdict"], "why": h["why"],
                          "tests": {t["id"]: {"status": t["status"], "value": t.get("value")} for t in h["tests"]},
                          "live_signals": [s["id"] for s in h["signals"] if s["live"]]} for h in sc_results],
          "companies": comp,
          "macro": {r.series: float(r.latest) for r in macro_changes.itertuples()}}
    path.write_text(json.dumps(js, indent=1, default=str))


def summary_md(path, asof, sc_results, scores, catalysts: pd.DataFrame, claims: pd.DataFrame, sent_trend,
               rel_sum: pd.DataFrame | None = None):
    L = [f"# Hypothesis tracker - {asof}", "",
         "_Research, not trading advice. Items marked **Inference** are interpretation, not data._", ""]
    L.append("## Verdicts")
    for h in sc_results:
        ch = ("first run, no earlier snapshot" if not h["prev_verdict"] else
              "unchanged" if h["prev_verdict"] == h["verdict"] else f"changed from **{h['prev_verdict']}**")
        L.append(f"- **{h['id']} {h['name']}: {h['verdict']}** ({h['why']}; {ch} since {h['prev_date'] or '-'})")
        for t in h["tests"]:
            v = t.get("value")
            vs = f"{v:+.3g} vs {t['op']} {t['threshold']}" if v is not None else ""
            L.append(f"  - {t['id']} {t['status'].upper()}: {t['name']} {vs}")
        live = [s for s in h["signals"] if s["live"]]
        L.append("  - Live signals: " + (", ".join(f"{s['kind']} {s['id']} {s['name']}" for s in live) or "none auto-detected")
                 + "; manual checks: " + ", ".join(s["id"] for s in h["signals"] if s["status"] == "manual"))
    L += ["", "## Three biggest sentiment-vs-fundamentals divergences (Inference)"]
    d = scores.dropna(subset=["gap"]).reindex(scores["gap"].abs().sort_values(ascending=False).index).dropna(subset=["gap"]).head(3)
    for tk, r in d.iterrows():
        kind = "fundamentals ahead of sentiment" if r.gap > 0 else "sentiment ahead of fundamentals"
        L.append(f"- **{r.company} ({tk})**: {kind}; fundamentals z {r.fundamentals_score:+.2f}, "
                 f"sentiment z {r.sentiment_score:+.2f}.")
    if sent_trend["change"].notna().any():
        st = sent_trend.dropna(subset=["change"])
        up, dn = st.loc[st.change.idxmax()], st.loc[st.change.idxmin()]
        L.append(f"- Sentiment moves since {st.prev_snapshot.iloc[0]}: biggest up {up.company} {up.change:+.2f}, "
                 f"biggest down {dn.company} {dn.change:+.2f}.")
    if rel_sum is not None and len(rel_sum):
        L += ["", "## Supplier-customer links"]
        for _, r in rel_sum.iterrows():
            line = (f"- **{r.supplier} -> {r.customer}** ({r.get('type', 'supplier')}): 1y correlation {r.corr_1y:.2f}, beta {r.beta_1y:.2f}")
            if pd.notna(r.get("last_event")):
                line += (f"; around {r.customer}'s {r.last_event} results {r.customer} moved {r.last_customer_move:+.1%}, "
                         f"{r.supplier} {r.last_supplier_move:+.1%} (same direction in "
                         f"{r.same_direction_share:.0%} of the last {int(r.earnings_events)} releases)")
            L.append(line)
    L += ["", "## Catalysts in the next 30 days"]
    for _, r in catalysts.iterrows():
        L.append(f"- {r.date}: {r.event}")
    diff = claims[claims.status == "differs"]
    if len(diff):
        L += ["", "## Newsletter figures that differ from the data"]
        for _, r in diff.iterrows():
            a = r.actual
            shown = f"{a:+.1%}" if r.check_type in ("move_in_window", "latest_quarter_yoy", "series_yoy", "fund_change", "price_vs_prior_high") else f"{a:.2f}"
            L.append(f'- {r.subject}: "{r.claim}". Data: {shown} ({r.detail})')
    L += ["", "Full detail: `tracker.xlsx`; charts in `charts/`."]
    path.write_text("\n".join(L) + "\n")
